"""Build the editable thesis from the Markdown parts and the faculty layout rules."""

from __future__ import annotations

import json
import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).parent
PARTS = [
    "00a-streszczenie.md", "00-wstep.md", "01-ai.md",
    "02-personalizacja.md", "03-zastosowania.md", "04-uwarunkowania.md",
    "05-przypadki.md", "06-zakonczenie.md", "07-bibliografia.md",
    "08-spis-tabel.md", "09-abstract.md",
]


def get_style(doc, name):
    try:
        return doc.styles[name]
    except KeyError:
        for style in doc.styles:
            if style.name == name:
                return style
        return doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


def prevent_hyphenation(style):
    prop = style.element.get_or_add_pPr()
    if prop.find(qn("w:suppressAutoHyphens")) is None:
        item = OxmlElement("w:suppressAutoHyphens")
        item.set(qn("w:val"), "1")
        prop.append(item)


def force_times(font_owner):
    font_owner.font.name = "Times New Roman"
    rpr = font_owner._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ["asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"]:
        fonts.attrib.pop(qn("w:" + attr), None)
    for attr in ["ascii", "hAnsi", "eastAsia", "cs"]:
        fonts.set(qn("w:" + attr), "Times New Roman")


def style_document(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.gutter = Cm(0.5)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.different_first_page_header_footer = False
    settings = doc.settings.element
    if settings.find(qn("w:mirrorMargins")) is None:
        settings.append(OxmlElement("w:mirrorMargins"))
    if settings.find(qn("w:doNotHyphenateCaps")) is None:
        settings.append(OxmlElement("w:doNotHyphenateCaps"))
    sect_props = sec._sectPr
    num = sect_props.find(qn("w:pgNumType"))
    if num is None:
        num = OxmlElement("w:pgNumType")
        sect_props.append(num)
    num.set(qn("w:start"), "3")

    for name in ["Normal", "Body Text", "First Paragraph"]:
        st = get_style(doc, name)
        st.font.name, st.font.size = "Times New Roman", Pt(12)
        fmt = st.paragraph_format
        fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        fmt.first_line_indent = Cm(0.75)
        fmt.line_spacing = 1.5
        fmt.space_before = fmt.space_after = Pt(0)
        fmt.widow_control = True
        prevent_hyphenation(st)
    for name, size, align, top, bottom, first_page in [
        ("Heading 1", 14, WD_ALIGN_PARAGRAPH.CENTER, 0, 12, True),
        ("Heading 2", 12, WD_ALIGN_PARAGRAPH.LEFT, 12, 12, False),
    ]:
        st = get_style(doc, name)
        st.font.name, st.font.size, st.font.bold = "Times New Roman", Pt(size), True
        fmt = st.paragraph_format
        fmt.alignment, fmt.first_line_indent = align, Cm(0)
        fmt.space_before, fmt.space_after = Pt(top), Pt(bottom)
        fmt.keep_with_next = fmt.keep_together = True
        fmt.page_break_before = first_page
        prevent_hyphenation(st)
    extra = [
        ("TOC Heading", 14, True, WD_ALIGN_PARAGRAPH.CENTER),
        ("TOC Entry", 10, False, WD_ALIGN_PARAGRAPH.LEFT),
        ("Table Caption", 10, False, WD_ALIGN_PARAGRAPH.CENTER),
        ("Table Source", 8, False, WD_ALIGN_PARAGRAPH.LEFT),
    ]
    for name, size, bold, align in extra:
        st = get_style(doc, name)
        st.font.name, st.font.size, st.font.bold = "Times New Roman", Pt(size), bold
        fmt = st.paragraph_format
        fmt.alignment, fmt.first_line_indent = align, Cm(0)
        fmt.space_before = fmt.space_after = Pt(0)
        fmt.line_spacing = 1 if name != "TOC Heading" else 1.5
        fmt.keep_with_next = name in ("TOC Heading", "Table Caption")
        prevent_hyphenation(st)
    for name in ["Normal", "Body Text", "First Paragraph", "Heading 1", "Heading 2",
                 "TOC Heading", "TOC Entry", "Table Caption", "Table Source"]:
        get_style(doc, name).font.color.rgb = RGBColor(0, 0, 0)
        force_times(get_style(doc, name))
    get_style(doc, "Table Caption").paragraph_format.space_before = Pt(12)
    get_style(doc, "Table Source").paragraph_format.space_after = Pt(12)
    get_style(doc, "Table Source").font.italic = True

    sec.footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer = sec.footer.paragraphs[0]
    footer.text = ""
    footer.style = get_style(doc, "Normal")
    footer.paragraph_format.first_line_indent = Cm(0)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)


def headings(markdown):
    output = []
    in_bibliography = False
    for line in markdown.splitlines():
        if line.startswith("# "):
            title = line[2:].strip()
            in_bibliography = title == "BIBLIOGRAFIA"
            output.append((title, 1))
        elif line.startswith("## ") and not in_bibliography:
            output.append((line[3:].strip(), 2))
    return output


def make_document():
    markdown = "\n\n".join((ROOT / "tekst" / name).read_text().strip() for name in PARTS) + "\n"
    (ROOT / "Praca_licencjacka_Bernaciak.md").write_text(markdown)
    body_file = ROOT / "tresc.docx"
    subprocess.run(["pandoc", str(ROOT / "Praca_licencjacka_Bernaciak.md"),
                    "-o", str(body_file)], check=True)
    doc = Document(body_file)
    style_document(doc)
    pages = json.loads((ROOT / "mapa-stron.json").read_text()) if (ROOT / "mapa-stron.json").exists() else {}
    toc_nodes = []
    p = doc.add_paragraph("SPIS TREŚCI", style=get_style(doc, "TOC Heading"))
    p.paragraph_format.space_after = Pt(12)
    toc_nodes.append(p._p)
    for title, level in headings(markdown):
        p = doc.add_paragraph(style=get_style(doc, "TOC Entry"))
        p.paragraph_format.left_indent = Cm(0.5 if level == 2 else 0)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(15.2), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        p.add_run(title)
        p.add_run("\t")
        p.add_run(str(pages.get(title, "–")))
        toc_nodes.append(p._p)
    for element in toc_nodes:
        doc.element.body.remove(element)
    for index, element in enumerate(toc_nodes):
        doc.element.body.insert(index, element)
    for p in doc.paragraphs:
        for run in p.runs:
            force_times(run)
        if re.match(r"^Tabela \d+\. ", p.text) and "— s." not in p.text:
            p.style = get_style(doc, "Table Caption")
        elif p.text.startswith("Źródło:"):
            p.style = get_style(doc, "Table Source")
        elif p.text.startswith("Tabela ") and " — s. " in p.text:
            p.style = get_style(doc, "TOC Entry")
        if re.search(r" {2,}", p.text):
            raise ValueError("Podwójna spacja w akapicie: " + p.text[:80])
    table_widths = [
        [2.5, 3.5, 4.5, 5.0],
        [3.0, 3.8, 4.5, 4.2],
        [2.5, 4.1, 4.1, 4.8],
        [7.5, 3.0, 5.0],
        [2.6, 3.2, 5.1, 4.6],
    ]
    for table_num, table in enumerate(doc.tables):
        table.autofit = False
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for col, width in zip(table.columns, table_widths[table_num]):
            col.width = Cm(width)
        borders = table._tbl.tblPr.find(qn("w:tblBorders"))
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            table._tbl.tblPr.append(borders)
        for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            tag = OxmlElement("w:" + edge)
            tag.set(qn("w:val"), "single")
            tag.set(qn("w:sz"), "4")
            tag.set(qn("w:color"), "888888")
            borders.append(tag)
        for row_num, row in enumerate(table.rows):
            row_props = row._tr.get_or_add_trPr()
            if row_props.find(qn("w:cantSplit")) is None:
                row_props.append(OxmlElement("w:cantSplit"))
            for cell_num, cell in enumerate(row.cells):
                cell.width = Cm(table_widths[table_num][cell_num])
                for p in cell.paragraphs:
                    p.paragraph_format.first_line_indent = Cm(0)
                    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_num == 0 else WD_ALIGN_PARAGRAPH.LEFT
                    for run in p.runs:
                        run.font.name, run.font.size = "Times New Roman", Pt(10)
                        run.bold = row_num == 0
        trpr = table.rows[0]._tr.get_or_add_trPr()
        if trpr.find(qn("w:tblHeader")) is None:
            trpr.append(OxmlElement("w:tblHeader"))
    doc.core_properties.title = "Zastosowanie sztucznej inteligencji do personalizacji relacji z klientami sklepów internetowych"
    doc.core_properties.author = "Łukasz Bernaciak"
    doc.core_properties.subject = "Praca licencjacka – wersja do uzupełnienia o stronę tytułową i oświadczenie"
    doc.save(ROOT / "Praca_licencjacka_Bernaciak.docx")
    body_file.unlink()
    print("DOCX zapisany; tabel:", len(doc.tables), "; pozycji spisu treści:", len(headings(markdown)))


if __name__ == "__main__":
    make_document()
