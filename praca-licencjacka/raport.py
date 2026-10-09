"""Verify the edited thesis and refresh the report and download package."""

from pathlib import Path
import json
import re
import zipfile

import fitz
from docx import Document
from docx.oxml.ns import qn

root = Path(__file__).parent
pdf = fitz.open(root / "Praca_licencjacka_Bernaciak.pdf")
doc = Document(root / "Praca_licencjacka_Bernaciak.docx")
markdown = (root / "Praca_licencjacka_Bernaciak.md").read_text()
stats = json.loads((root / "statystyki.json").read_text())
sources = json.loads((root / "zrodla-rejestr.json").read_text())
page_map = json.loads((root / "mapa-stron.json").read_text())
page_texts = [re.sub(r"\s+", " ", page.get_text()) for page in pdf]

section = doc.sections[0]
assert len(pdf) == 46 and len(pdf) + 2 == 48
assert round(section.page_width.cm, 1) == 21 and round(section.page_height.cm, 1) == 29.7
assert all(round(getattr(section, field).cm, 1) == value for field, value in [
    ("left_margin", 2.5), ("right_margin", 2.5), ("gutter", 0.5),
    ("top_margin", 2.5), ("bottom_margin", 2.5),
])
assert doc.settings.element.find(qn("w:mirrorMargins")) is not None
assert section._sectPr.find(qn("w:pgNumType")).get(qn("w:start")) == "3"
for name, size in [("Normal", 12), ("Heading 1", 14), ("Heading 2", 12),
                   ("Table Caption", 10), ("Table Source", 8)]:
    style = next(x for x in doc.styles if x.name == name)
    assert style.font.name == "Times New Roman" and style.font.size.pt == size
assert round(next(x for x in doc.styles if x.name == "Normal").paragraph_format.first_line_indent.cm, 2) == 0.75
assert next(x for x in doc.styles if x.name == "Normal").paragraph_format.line_spacing == 1.5
assert len(doc.tables) == 5 and len(sources) == 26
assert 12 <= stats["total"]["mean"] <= 15

polish = (root / "tekst/00a-streszczenie.md").read_text().split("\n\n")[1]
summary_chars = len(re.sub(r"\s", "", polish))
assert 1000 <= summary_chars <= 1500
assert "Keywords:" in (root / "tekst/09-abstract.md").read_text()
assert not any(re.search(r" {2,}", p.text) for p in doc.paragraphs)
assert not any(re.search(r" {2,}", c.text) for t in doc.tables for row in t.rows for c in row.cells)
assert not any(re.search(r" {2,}", p.read_text()) for p in (root / "tekst").glob("*.md"))
cut_lines = [(page_no + 3, line) for page_no, page in enumerate(pdf)
             for line in page.get_text().splitlines() if line.rstrip().endswith("-")]
assert not cut_lines, cut_lines

for title, printed_page in page_map.items():
    assert title in page_texts[printed_page - 3], (title, printed_page)
table_list = (root / "tekst/08-spis-tabel.md").read_text()
for number in range(1, 6):
    actual = next(i + 3 for i, text in enumerate(page_texts)
                  if i > 0 and re.search(rf"Tabela {number}\. ", text))
    line = next(line for line in table_list.splitlines() if line.startswith(f"Tabela {number}. "))
    assert line.endswith(f"— s. {actual}"), line

citations = [
    "OECD 2019", "Adam et al. 2021", "Adomavicius i Tuzhilin 2005",
    "Aguirre et al. 2015", "Bleier i Eisenbeiss 2015", "Huang i Rust 2021",
    "LeCun et al. 2015", "Lemon i Verhoef 2016", "Linden et al. 2003",
    "Luo et al. 2019", "Payne i Frow 2005", "Abadi et al. 2016",
    "Jouppi et al. 2017", "Paszke et al. 2019", "Vaswani et al. 2017",
    "Dz.Urz. UE L 201 z 2002", "Dz.Urz. UE L 304 z 2011",
    "Dz.Urz. UE L 119 z 2016", "Dz.Urz. UE L z 2024",
    "Dz.U. z 2024 r., poz. 1221", "www.aboutamazon.com/news/retail/amazon-rufus",
    "blog.allegro.tech/2026/08/automating-search-relevance", "McCarthy 2007",
    "Mell i Grance 2011", "NIST 2023",
    "corporate.zalando.com/en/technology/zalando-launch",
]
content = "\n".join(p.read_text() for p in (root / "tekst").glob("*.md")
                    if p.name not in ("07-bibliografia.md", "08-spis-tabel.md"))
assert len(citations) == len(sources) and all(citation in content for citation in citations)

report = f"""# Raport weryfikacji redakcyjnej

## Wynik

Przygotowano poprawioną pracę licencjacką na podstawie dotychczasowego tekstu oraz załączonego szablonu WZNJ. Zgodnie z decyzją autora pliki nie zawierają strony tytułowej ani oświadczenia; autor dołączy obie strony samodzielnie.

- PDF: **{len(pdf)} stron pliku**, numerowanych od **3 do {len(pdf)+2}**. Po dołączeniu dwóch stron początkowych całość będzie miała **{len(pdf)+2} stron**.
- Bibliografia: **{len(sources)} pozycji**, wszystkie przywołane w tekście.
- Tabele: **{len(doc.tables)}**, z numeracją ciągłą i sprawdzonymi stronami w spisie.
- Streszczenie polskie: **{summary_chars} znaków bez spacji**; pięć słów kluczowych.
- Średnia długość zdania tekstu głównego: **{stats['total']['mean']:.2f} słowa**.
- Podwójne spacje w tekście DOCX i Markdown: **0**.
- Wyrazy lub adresy zakończone dywizem na końcu wiersza PDF: **0**.

## Zastosowane wymagania

Dokument ma format A4, marginesy lustrzane po 2,5 cm i 0,5 cm na oprawę. Numer strony znajduje się na środku stopki. Tekst główny używa stylu Times New Roman 12 pkt, interlinii 1,5, justowania, wcięcia 0,75 cm i zerowych odstępów między akapitami. Tytuły rozdziałów mają 14 pkt i są wyśrodkowane; tytuły podrozdziałów mają 12 pkt i są wyrównane do lewej. Główne części zaczynają się od nowej strony. Wyłączono automatyczne dzielenie wyrazów.

Tabele są wyśrodkowane, numerowane kolejno od 1 do 5 i mają tekst 10 pkt z pojedynczą interlinią. Tytuły o wielkości 10 pkt znajdują się nad tabelami, a źródła pod nimi. Bibliografię ułożono według rodzajów źródeł wymaganych w załącznikach i ponumerowano. W tekście zastosowano wskazany w szablonie zapis odwołań. Dodano polskie streszczenie i słowa kluczowe oraz angielski tytuł, streszczenie i słowa kluczowe.

## Sprawdzenia i ograniczenia

Sprawdzono strukturę DOCX, strony sekcji, widoczne tabele, spis treści, spis tabel, numerację oraz tekst wyodrębniony z PDF. PDF został wyeksportowany z DOCX w LibreOffice, dlatego oba pliki mają zgodny skład. W środowisku generującym PDF nie ma zainstalowanego kroju Times New Roman; eksport użył metrycznie zbliżonego **Liberation Serif**. DOCX zachowuje ustawienie Times New Roman. Aby otrzymać PDF z rzeczywistym Times New Roman, należy otworzyć DOCX na komputerze z licencjonowaną czcionką i wyeksportować go ponownie.

Stronę tytułową, oświadczenie, wymagane podpisy i zgody autor dołączy samodzielnie. Fiszka bibliograficzna również pozostaje do uzupełnienia o dane promotora i decyzje autora. Raport dotyczy redakcji i składu; wcześniejsze ograniczenia weryfikacji pełnych tekstów publikacji oraz aktualności aktów prawnych pozostają aktualne.
"""
report += "\n" + (root / "Weryfikacja_zrodel.md").read_text().rstrip() + "\n"
(root / "Raport_weryfikacji.md").write_text(report)

bundle = ["Praca_licencjacka_Bernaciak.docx", "Praca_licencjacka_Bernaciak.pdf",
          "Praca_licencjacka_Bernaciak.md", "Raport_weryfikacji.md",
          "statystyki.json", "zrodla-rejestr.json", "mapa-stron.json", "Weryfikacja_zrodel.md",
          "dokument.py", "bibliografia.py", "statystyki.py", "raport.py"]
bundle += ["tekst/" + p.name for p in sorted((root / "tekst").glob("*.md"))]
with zipfile.ZipFile(root / "Praca_licencjacka_Bernaciak_pakiet.zip", "w", zipfile.ZIP_DEFLATED) as archive:
    for name in bundle:
        archive.write(root / name, name)
with zipfile.ZipFile(root / "Praca_licencjacka_Bernaciak_pakiet.zip") as archive:
    assert archive.testzip() is None
print(json.dumps({"pdf_pages": len(pdf), "printed_last_page": len(pdf) + 2,
                  "sources": len(sources), "tables": len(doc.tables),
                  "summary_chars": summary_chars, "package_files": len(bundle)}, ensure_ascii=False))
